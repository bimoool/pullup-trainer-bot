from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.db.models import Baseline
from app.domain.constants import EquipmentType, ExerciseType
from app.services.export import build_export_workbook
from tests.test_reports import _record


def test_build_export_workbook_creates_all_four_sheets():
    records = [_record(1), _record(4, a_equipment_type=EquipmentType.BODYWEIGHT, a_equipment_value=None)]
    baselines = [Baseline(performed_at=datetime(2025, 12, 1, tzinfo=UTC), reps=8)]

    buffer = build_export_workbook(records, baselines, set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    assert workbook.sheetnames == ["История", "Замеры", "Прогресс", "Продолжить офлайн"]


def test_history_sheet_has_header_and_one_row_per_workout():
    records = [_record(1), _record(4)]
    buffer = build_export_workbook(records, baselines=[], set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["История"]
    assert ws.max_row == 1 + len(records)
    assert ws.cell(row=1, column=1).value == "Дата"
    assert ws.cell(row=2, column=1).value == "01.01.2026"
    # снаряд блока на объём (после колонок Цикл/Направление) — без лишних
    # нулей в кг (Часть 10: 20.0 -> 20).
    assert ws.cell(row=2, column=4).value == "резина 20 кг"


def test_history_sheet_splits_working_reps_into_per_set_columns():
    # _record даёт a_reps=(15,15,15), a_max=16 по умолчанию (VOLUME_BLOCK.work_sets=3)
    records = [_record(1)]
    buffer = build_export_workbook(records, baselines=[], set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["История"]
    header = [cell.value for cell in ws[1]]
    assert "Объём: подход 1" in header
    assert "Объём: подход 3" in header
    assert "Объём: максимум" in header
    volume_set_1_col = header.index("Объём: подход 1") + 1
    volume_max_col = header.index("Объём: максимум") + 1
    assert ws.cell(row=2, column=volume_set_1_col).value == 15
    assert ws.cell(row=2, column=volume_max_col).value == 16


def test_history_sheet_shows_cycle_number_and_exercise_type():
    records = [_record(1, workout_set_id=42, exercise_type=ExerciseType.PULL_UPS)]
    buffer = build_export_workbook(records, baselines=[], set_numbers={42: 3})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["История"]
    assert ws.cell(row=1, column=2).value == "Цикл"
    assert ws.cell(row=2, column=2).value == 3  # реальный set_number, не позиция в списке
    assert ws.cell(row=1, column=3).value == "Направление"
    assert ws.cell(row=2, column=3).value == "подтягивания"


def test_history_sheet_blank_cycle_when_workout_set_id_unknown():
    records = [_record(1)]  # workout_set_id=None по умолчанию
    buffer = build_export_workbook(records, baselines=[], set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["История"]
    # openpyxl округляет пустую строку до None при обратном чтении из файла —
    # для пользователя ячейка всё равно выглядит просто пустой.
    assert ws.cell(row=2, column=2).value in (None, "")


def test_baselines_sheet_lists_each_baseline():
    baselines = [
        Baseline(performed_at=datetime(2025, 12, 1, tzinfo=UTC), reps=8),
        Baseline(performed_at=datetime(2026, 2, 1, tzinfo=UTC), reps=12),
    ]
    buffer = build_export_workbook(records=[], baselines=baselines, set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["Замеры"]
    assert ws.max_row == 1 + len(baselines)
    assert ws.cell(row=2, column=2).value == 8
    assert ws.cell(row=3, column=2).value == 12


def test_progress_sheet_groups_by_iso_week():
    same_week = [_record(1), _record(3)]  # 1 и 3 января 2026 — один календарный год/неделя
    buffer = build_export_workbook(records=same_week, baselines=[], set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["Прогресс"]
    assert ws.max_row == 2  # заголовок + одна неделя
    assert ws.cell(row=2, column=2).value == 2  # 2 тренировки в этой неделе


def test_offline_template_sheet_has_no_data_but_ready_sum_formulas():
    # Часть 10: шаблон по-прежнему без данных (заполняет человек офлайн),
    # но объём теперь считает сам Excel — заранее подготовленные строки с
    # формулой =SUM(...), не руками.
    buffer = build_export_workbook(records=[_record(1)], baselines=[], set_numbers={})
    workbook = load_workbook(BytesIO(buffer.read()))

    ws = workbook["Продолжить офлайн"]
    assert ws.cell(row=1, column=1).value == "Дата"
    header = [cell.value for cell in ws[1]]
    assert "Объём: сумма" in header
    assert "Сила: сумма" in header

    volume_sum_col = header.index("Объём: сумма") + 1
    strength_sum_col = header.index("Сила: сумма") + 1
    volume_formula = ws.cell(row=2, column=volume_sum_col).value
    strength_formula = ws.cell(row=2, column=strength_sum_col).value
    assert volume_formula.startswith("=SUM(")
    assert strength_formula.startswith("=SUM(")
    # ни одно из полей ввода не заполнено — это шаблон, не готовые данные
    assert ws.cell(row=2, column=2).value is None
