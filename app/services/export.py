import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.bot.formatting import format_kg
from app.db.models import Baseline
from app.domain.constants import STRENGTH_BLOCK, VOLUME_BLOCK, EquipmentType, ExerciseType
from app.domain.session import BlockAssignment, WorkoutRecord

_EQUIPMENT_LABELS = {
    EquipmentType.BAND: "резина",
    EquipmentType.BODYWEIGHT: "свой вес",
    EquipmentType.WEIGHT: "отягощение",
    EquipmentType.AUSTRALIAN: "австралийские",
}

# Сейчас всегда PULL_UPS, но колонка в истории уже общая — задел под
# будущие направления тренировок (Часть 8 респека), не только подтягивания.
_EXERCISE_LABELS = {
    ExerciseType.PULL_UPS: "подтягивания",
}

# Часть 10: вместо одного столбца "макс" — отдельный столбец на каждый
# рабочий подход + отдельно максимум, так видно всю тренировку по числам,
# не только финальный результат. Число столбцов берётся из домена
# (VOLUME_BLOCK.work_sets/STRENGTH_BLOCK.work_sets), не захардкожено.
_VOLUME_SET_HEADERS = [f"Объём: подход {i + 1}" for i in range(VOLUME_BLOCK.work_sets)]
_STRENGTH_SET_HEADERS = [f"Сила: подход {i + 1}" for i in range(STRENGTH_BLOCK.work_sets)]

_HISTORY_HEADER = [
    "Дата", "Цикл", "Направление",
    "Объём: снаряд", *_VOLUME_SET_HEADERS, "Объём: максимум", "Объём: цель",
    "Сила: снаряд", *_STRENGTH_SET_HEADERS, "Сила: максимум", "Сила: цель",
    "Комментарий",
]
_BASELINES_HEADER = ["Дата", "Повторения"]
_PROGRESS_HEADER = ["Неделя", "Тренировок", "Суммарный объём"]
_OFFLINE_HEADER = [
    "Дата",
    *_VOLUME_SET_HEADERS, "Объём: максимум", "Объём: сумма",
    *_STRENGTH_SET_HEADERS, "Сила: максимум", "Сила: сумма",
    "Комментарий",
]
# Сколько пустых строк с готовой формулой суммы заранее подготовить в
# офлайн-шаблоне (Часть 10) — чтобы объём считался Excel'ем (=SUM), а не
# руками, с первой же тренировки, вписанной уже без бота.
_OFFLINE_TEMPLATE_ROWS = 200

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_header_row(ws: Worksheet, column_count: int) -> None:
    for column in range(1, column_count + 1):
        cell = ws.cell(row=1, column=column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
    ws.freeze_panes = "A2"


def _auto_size_columns(ws: Worksheet, headers: list[str], *, min_width: int = 10, max_width: int = 22) -> None:
    for index, header in enumerate(headers, start=1):
        width = max(min_width, min(max_width, len(header) + 2))
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width


def _equipment_label(block: BlockAssignment) -> str:
    label = _EQUIPMENT_LABELS[block.equipment_type]
    if block.equipment_value is not None:
        return f"{label} {format_kg(block.equipment_value)} кг"
    return label


def _padded_working_reps(block: BlockAssignment, expected_sets: int) -> list[int | str]:
    """Рабочих подходов должно быть ровно expected_sets, но на всякий
    случай не падаем на рассинхроне (старые данные и т.п.) — недостающее
    остаётся пустым, лишнее не выводится."""
    reps = list(block.log.working_reps[:expected_sets])
    reps += [""] * (expected_sets - len(reps))
    return reps


def _write_history_sheet(ws: Worksheet, records: list[WorkoutRecord], set_numbers: dict[int, int]) -> None:
    """set_numbers — workout_set_id -> реальный порядковый номер цикла
    (WorkoutSet.set_number), не позиция в списке: закрытый досрочно цикл
    (см. "Завершить цикл") не должен сдвигать нумерацию следующих."""
    ws.append(_HISTORY_HEADER)
    for record in records:
        cycle = set_numbers.get(record.workout_set_id, "") if record.workout_set_id is not None else ""
        exercise_label = _EXERCISE_LABELS.get(record.exercise_type, record.exercise_type) if record.exercise_type else ""
        ws.append([
            record.performed_at.strftime("%d.%m.%Y"), cycle, exercise_label,
            _equipment_label(record.block_a), *_padded_working_reps(record.block_a, VOLUME_BLOCK.work_sets),
            record.block_a.log.max_reps, record.block_a.target_after,
            _equipment_label(record.block_b), *_padded_working_reps(record.block_b, STRENGTH_BLOCK.work_sets),
            record.block_b.log.max_reps, record.block_b.target_after,
            record.comment or "",
        ])
    _style_header_row(ws, len(_HISTORY_HEADER))
    _auto_size_columns(ws, _HISTORY_HEADER)


def _write_baselines_sheet(ws: Worksheet, baselines: list[Baseline]) -> None:
    ws.append(_BASELINES_HEADER)
    for baseline in baselines:
        ws.append([baseline.performed_at.strftime("%d.%m.%Y"), baseline.reps])
    _style_header_row(ws, len(_BASELINES_HEADER))
    _auto_size_columns(ws, _BASELINES_HEADER)


def _write_progress_sheet(ws: Worksheet, records: list[WorkoutRecord]) -> None:
    """Объём по неделям (ISO-неделя даты тренировки) — та же метрика, что
    в еженедельном отчёте (app/domain/reports.weekly_summary), но за всю
    историю целиком, а не только за последнюю неделю."""
    ws.append(_PROGRESS_HEADER)
    weekly: dict[tuple[int, int], list[WorkoutRecord]] = defaultdict(list)
    for record in records:
        iso = record.performed_at.isocalendar()
        weekly[(iso[0], iso[1])].append(record)

    for (year, week), week_records in sorted(weekly.items()):
        total_volume = sum(r.block_a.log.volume + r.block_b.log.volume for r in week_records)
        ws.append([f"{year}-W{week:02d}", len(week_records), total_volume])
    _style_header_row(ws, len(_PROGRESS_HEADER))
    _auto_size_columns(ws, _PROGRESS_HEADER)


def _write_offline_template_sheet(ws: Worksheet) -> None:
    """Шаблон для тех, кто боится потерять записи после окончания подписки
    (см. Часть 5 респека) — без данных, но с готовыми формулами суммы
    (Часть 10): объём считает сам Excel (=SUM), не человек руками. Разбивка
    по подходам — та же, что в "Истории"."""
    ws.append(_OFFLINE_HEADER)

    volume_set_cols = range(2, 2 + VOLUME_BLOCK.work_sets)  # после "Дата"
    volume_max_col = 2 + VOLUME_BLOCK.work_sets
    volume_sum_col = volume_max_col + 1
    strength_set_cols = range(volume_sum_col + 1, volume_sum_col + 1 + STRENGTH_BLOCK.work_sets)
    strength_max_col = volume_sum_col + 1 + STRENGTH_BLOCK.work_sets
    strength_sum_col = strength_max_col + 1

    def _col_letter(col: int) -> str:
        return ws.cell(row=1, column=col).column_letter

    for row in range(2, 2 + _OFFLINE_TEMPLATE_ROWS):
        volume_range = f"{_col_letter(volume_set_cols.start)}{row}:{_col_letter(volume_max_col)}{row}"
        strength_range = f"{_col_letter(strength_set_cols.start)}{row}:{_col_letter(strength_max_col)}{row}"
        ws.cell(row=row, column=volume_sum_col, value=f"=SUM({volume_range})")
        ws.cell(row=row, column=strength_sum_col, value=f"=SUM({strength_range})")

    _style_header_row(ws, len(_OFFLINE_HEADER))
    _auto_size_columns(ws, _OFFLINE_HEADER)


def build_export_workbook(
    records: list[WorkoutRecord], baselines: list[Baseline], set_numbers: dict[int, int],
) -> io.BytesIO:
    workbook = Workbook()
    history_sheet = workbook.active
    history_sheet.title = "История"
    _write_history_sheet(history_sheet, records, set_numbers)

    _write_baselines_sheet(workbook.create_sheet("Замеры"), baselines)
    _write_progress_sheet(workbook.create_sheet("Прогресс"), records)
    _write_offline_template_sheet(workbook.create_sheet("Продолжить офлайн"))

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
